#!/usr/bin/env python3
"""Build the LinkedIn post archive.

Run it from the author's working folder (or pass --root): drafts live there as *.md,
the archive is written to <root>/archive/. Nothing is read or written outside --root
except the export files you pass explicitly.

Sources
  1. local drafts   <root>/*.md
  2. LinkedIn export Shares.csv  (--shares path)  full text of every published post
  3. AggregateAnalytics XLSX     (--metrics path) reach / engagement, author-only data
  4. LinkedIn export Rich_Media.csv (--media path) post format, joined on the timestamp

Output
  archive/posts.jsonl   one JSON object per post, the machine-readable source of truth
  archive/ANGLES.md     human/agent index: topic x angle x date x reach

Manual fields (angle, rubric, topic_tags, notes) are preserved across rebuilds:
the script merges onto the existing jsonl instead of overwriting it.
"""

import argparse
import csv
import json
import re
import sys
from pathlib import Path
from urllib.parse import unquote

# set in main() from --root (default: current working directory)
ROOT = Path.cwd()
ARCHIVE = ROOT / "archive"
JSONL = ARCHIVE / "posts.jsonl"
INDEX = ARCHIVE / "ANGLES.md"

MANUAL_FIELDS = (
    "angle",
    "rubric",
    "topic_tags",
    "notes",
    "post_type",
    "format",
    "media_format",
    "funnel_stage",
)
DATE_IN_NAME = re.compile(r"(\d{4})-(\d{2})-(\d{2})")
SHARE_ID = re.compile(r"(?:share|ugcPost)[-:](\d+)")


def load_existing():
    """Existing records keyed by id, so manual classification survives a rebuild."""
    if not JSONL.exists():
        return {}
    out = {}
    for line in JSONL.read_text(encoding="utf-8").splitlines():
        if line.strip():
            rec = json.loads(line)
            out[rec["id"]] = rec
    return out


SKIP_LINE = re.compile(r"^(#|-{1,3}$|[-*>|`]|\*\*|\d+\.\s)")


def hook_of(text):
    """First line of actual post copy: skip headings, front-matter and note bullets."""
    for line in text.splitlines():
        line = line.strip()
        if line and not SKIP_LINE.match(line):
            return line
    return ""


def post_id(url, date, hook):
    m = SHARE_ID.search(url or "")
    if m:
        return m.group(1)
    return f"{date}-{re.sub(r'[^a-z0-9]+', '-', hook.lower())[:40]}".strip("-")


def from_local_md():
    """Drafts written in this repo. Some are published, some never were."""
    for path in sorted(ROOT.glob("*.md")):
        text = path.read_text(encoding="utf-8").strip()
        if not text:
            continue
        m = DATE_IN_NAME.search(path.name)
        date = "-".join(m.groups()) if m else ""
        # strip a leading "# Title" line: it is a file label, not post copy
        body = re.sub(r"^#\s+.*\n+", "", text)
        # drafts carry a brief above the first standalone "---"; it is notes, not post copy.
        # only cut when that separator sits near the top, so a mid-post rule is never eaten.
        head = body.split("\n")[:25]
        if "---" in [l.strip() for l in head]:
            cut = [l.strip() for l in head].index("---")
            body = "\n".join(body.split("\n")[cut + 1:]).strip()
        # some drafts carry a SECOND notes block (gate scores, references) before the copy
        NOTES = ("Гейт", "Хук взято", "Референс", "довга версія", "коротка версія", "Хард-фейл")
        if body.startswith(NOTES):
            lines = body.split("\n")
            marks = [i for i, l in enumerate(lines[:40]) if l.strip() == "---"]
            if marks:
                body = "\n".join(lines[marks[0] + 1:]).strip()
        hook = hook_of(body)
        yield {
            "id": post_id("", date, hook),
            "date": date,
            "url": "",
            "source": "local-md",
            "file": path.name,
            "hook": hook,
            "text": body,
            "reach": None,
            "engagement": None,
        }


QUOTE_BEFORE_NL = re.compile(r'"+(?=\n)')
QUOTE_AFTER_NL = re.compile(r'(?<=\n)"+')


def clean_export_text(text):
    """Shares.csv wraps every newline in stray double quotes: `line"\n""\n"line`.

    Strip only the quotes that touch a newline, keeping the newlines themselves so
    the blank-line rhythm of the post survives. A line that legitimately ended in a
    double quote loses it — rare enough to accept, and «» is the usual quote here.
    """
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = QUOTE_BEFORE_NL.sub("", text)
    text = QUOTE_AFTER_NL.sub("", text)
    return text.strip().strip('"').strip()


URN_IN_URL = re.compile(r"urn:li:(share|ugcPost|activity):(\d+)")


def open_url_of(url):
    """Clickable link. The export gives a percent-encoded urn; LinkedIn 307-redirects
    a share urn to the real activity, so the urn form is what actually opens."""
    m = URN_IN_URL.search(unquote(url or ""))
    if not m:
        return ""
    return f"https://www.linkedin.com/feed/update/urn:li:{m.group(1)}:{m.group(2)}/"


def from_shares_csv(path):
    """LinkedIn data export. Column names vary by export vintage, so detect them."""
    with open(path, newline="", encoding="utf-8-sig") as fh:
        reader = csv.DictReader(fh)
        cols = {c.lower().strip(): c for c in (reader.fieldnames or [])}

        def pick(*names):
            for n in names:
                if n in cols:
                    return cols[n]
            return None

        c_date = pick("date", "shareddate", "createdat")
        c_text = pick("sharecommentary", "commentary", "text", "content")
        c_link = pick("sharelink", "link", "url", "permalink")
        if not c_text:
            sys.exit(f"No text column in {path}. Columns: {reader.fieldnames}")

        for row in reader:
            text = clean_export_text(row.get(c_text) or "")
            if not text:
                continue  # reshare with no commentary of its own
            date = (row.get(c_date) or "")[:10]
            # the export percent-encodes the urn (urn%3Ali%3Ashare%3A123); decode it
            # or the share id never matches and every post gets a date-hook id instead
            url = unquote((row.get(c_link) or "").strip())
            hook = hook_of(text)
            yield {
                "id": post_id(url, date, hook),
                "date": date,
                # full timestamp, kept because Rich_Media.csv joins on it to the minute
                "posted_at": (row.get(c_date) or "").strip(),
                "url": url,
                "open_url": open_url_of(url),
                "source": "linkedin-export",
                "file": "",
                "hook": hook,
                "text": text,
                "reach": None,
                "engagement": None,
            }


def _norm(text):
    return re.sub(r"[^0-9a-zа-яёїієґ]+", "", (text or "").lower())


def collapse_drafts(records):
    """A draft and its published post are one post, under two ids: the draft has no
    URL so it falls back to a date-hook id, the published one carries the share id.

    Hooks alone don't match them — the first line usually gets rewritten before
    publishing — and neither does whole-text similarity, because a draft file also
    holds the creative brief around the copy. What survives both is a shared block
    of text, so score on the longest common run and keep the export record (real
    copy, real id, real metrics), carrying the draft's file reference and manual
    classification onto it.
    """
    from datetime import date
    from difflib import SequenceMatcher

    def as_date(value):
        try:
            return date(*(int(p) for p in value.split("-")))
        except (ValueError, TypeError, AttributeError):
            return None

    def same_opening(draft, post):
        """A scraped record keeps the published opening almost verbatim, even when its
        body was truncated — so a long shared hook prefix settles the pair on its own."""
        a, b = _norm(draft.get("hook")), _norm(post.get("hook"))
        if not a or not b:
            return False
        shared = 0
        for x, y in zip(a, b):
            if x != y:
                break
            shared += 1
        return shared >= HOOK_PREFIX_MIN

    def shared_run(draft_text, post_text):
        """Longest common run, as a share of the shorter text. Real pairs land around
        0.27-0.32 here and unrelated same-week posts under 0.03, so the gap is wide."""
        matcher = SequenceMatcher(None, draft_text, post_text, autojunk=False)
        block = matcher.find_longest_match(0, len(draft_text), 0, len(post_text))
        return block.size / max(1, min(len(draft_text), len(post_text)))

    published = [r for r in records.values() if r.get("source") == "linkedin-export"]
    dropped = 0

    others = [r for r in records.values() if r.get("source") in COLLAPSIBLE_SOURCES]
    for draft in others:
        d_date, d_text = as_date(draft.get("date")), _norm(draft.get("text"))
        if not d_date or not d_text:
            continue
        # a scraped date is the day it was read, not the day it was posted
        window = DRAFT_WINDOW_DAYS if draft["source"] == "local-md" else SCRAPE_WINDOW_DAYS
        best, best_score = None, 0.0
        for post in published:
            p_date = as_date(post.get("date"))
            if not p_date or abs((p_date - d_date).days) > window:
                continue
            score = shared_run(d_text, _norm(post.get("text")))
            if same_opening(draft, post):
                score = max(score, DRAFT_MATCH_MIN)
            if score > best_score:
                best, best_score = post, score
        if not best or best_score < DRAFT_MATCH_MIN:
            continue  # never published, or published too differently to claim a match
        best["file"] = best.get("file") or draft.get("file", "")
        best["draft_of"] = draft["id"]
        for field in MANUAL_FIELDS:
            if draft.get(field) and not best.get(field):
                best[field] = draft[field]
        # a scrape carries counts the export does not; never lose them to a merge
        for field in ("reach", "engagement"):
            if best.get(field) is None and draft.get(field) is not None:
                best[field] = draft[field]
        del records[draft["id"]]
        dropped += 1

    if dropped:
        print(f"drafts collapsed onto their published post: {dropped}")
    return records


COLLAPSIBLE_SOURCES = ("local-md", "linkedin-web")
DRAFT_WINDOW_DAYS = 3
SCRAPE_WINDOW_DAYS = 7
DRAFT_MATCH_MIN = 0.10
HOOK_PREFIX_MIN = 25

MEDIA_STAMP = re.compile(
    r"uploaded an? (?P<kind>[a-z ]+?) on (?P<month>\w+) (?P<day>\d+), (?P<year>\d{4})"
    r" at (?P<hour>\d+):(?P<minute>\d\d) (?P<meridiem>AM|PM)",
    re.IGNORECASE,
)
MONTHS = {
    m: i
    for i, m in enumerate(
        "january february march april may june july august september october "
        "november december".split(),
        start=1,
    )
}
# a LinkedIn carousel IS a multi-page document post, so "feed document" is the
# carousel signal; a one-page PDF lead magnet lands here too and needs a manual flip
MEDIA_FORMAT = {
    "feed document": "карусель",
    "document": "документ/PDF",
    "feed photo": "single image",
    "photo": "single image",
    "video": "відео",
}


def _as_minutes(stamp):
    """"YYYY-MM-DD HH:MM" → minutes since epoch-ish, for cheap window arithmetic."""
    from datetime import datetime

    try:
        moment = datetime.strptime(stamp[:16], "%Y-%m-%d %H:%M")
    except ValueError:
        return None
    return int(moment.timestamp() // 60)


def from_media_csv(path):
    """Rich_Media.csv from the LinkedIn export → [(minute, format)] sorted by time.

    The file is useless for text — LinkedIn replaces every Cyrillic character with
    "?" on its own side — but the upload kind and its timestamp survive intact, and
    that pairing is what tells a carousel from an infographic from a plain text post.
    """
    uploads = []
    with open(path, newline="", encoding="utf-8-sig") as fh:
        for row in csv.DictReader(fh):
            raw = " ".join(v for v in row.values() if v)
            m = MEDIA_STAMP.search(raw)
            if not m:
                continue
            kind = m.group("kind").strip().lower()
            if kind == "profile photo":
                continue  # avatar change, not a post
            hour = int(m.group("hour")) % 12
            if m.group("meridiem").upper() == "PM":
                hour += 12
            stamp = (
                f"{m.group('year')}-{MONTHS[m.group('month').lower()]:02d}"
                f"-{int(m.group('day')):02d} {hour:02d}:{m.group('minute')}"
            )
            minute = _as_minutes(stamp)
            if minute is not None:
                uploads.append((minute, MEDIA_FORMAT.get(kind, kind)))
    uploads.sort()
    print(f"media: {len(uploads)} uploads read from {Path(path).name}")
    return uploads


QUOTE = r"[«»\"“”„‘’]"
# "Напишіть «СКІЛИ» в коментарях" — the quoted keyword is the giveaway
GATE_KEYWORD = re.compile(
    r"(?:напиш\w*|прокоментуй\w*|пишіть)[^.!?\n]{0,40}" + QUOTE + r"\s*\w[^\n]{0,30}?" + QUOTE,
    re.IGNORECASE,
)
GATE_COMMENT = re.compile(r"коментар", re.IGNORECASE)
GATE_DELIVERY = re.compile(
    r"(?:надішлю|надішлем|скину|кину|відправлю|вишлю)[^.!?\n]{0,60}"
    r"(?:особист|дірект|директ|приват|inbox|дм)",
    re.IGNORECASE,
)


def is_lead_magnet(text):
    """A lead magnet is not a media type, it is a CTA: comment for the asset, get it
    in DM. So it is read off the copy — either an imperative with a quoted keyword,
    or a mention of comments paired with a promise to send something privately."""
    if GATE_KEYWORD.search(text or ""):
        return True
    return bool(GATE_COMMENT.search(text or "") and GATE_DELIVERY.search(text or ""))


def apply_lead_magnets(records):
    """Lead magnet outranks the media type: what matters is that the post trades an
    asset for a comment, not whether that asset went out as an image or a document."""
    auto = set(MEDIA_FORMAT.values()) | {"text-only", ""}
    tagged = 0
    for rec in records.values():
        current = rec.get("format") or ""
        # never overwrite a format set by hand — only one this script derived itself
        if current not in auto or not is_lead_magnet(rec.get("text")):
            continue
        rec["media_format"] = current
        rec["format"] = "лід-магніт"
        tagged += 1
    print(f"lead magnets: {tagged} (media type kept in `media_format`)")
    return records


def apply_media(records, uploads):
    """Set `format` from the media export, without overwriting a manual call.

    An upload lands before its post, not on it: a photo by the same minute, but a
    document by up to a quarter of an hour, because LinkedIn renders every page
    before the post can go out. So take the nearest upload inside a window ending
    at the post.

    Rich_Media.csv is INCOMPLETE — it logs roughly a fifth of the posts that carried
    media, in every year. A post missing from it is therefore UNKNOWN, never
    text-only: absence in a partial export is not evidence of absence of media.
    Unmatched posts keep an empty `format` and wait for a human.
    """
    filled, unknown = 0, 0
    for rec in records.values():
        if rec.get("format"):
            continue
        posted = _as_minutes(rec.get("posted_at") or "")
        candidates = (
            [
                (posted - minute, fmt)
                for minute, fmt in uploads
                if 0 <= posted - minute <= MEDIA_WINDOW_MIN
            ]
            if posted is not None
            else []
        )
        if candidates:
            rec["format"] = min(candidates)[1]
            filled += 1
        else:
            unknown += 1
    print(f"format set: {filled} from media log, {unknown} left unknown (not text-only)")
    return records


MEDIA_WINDOW_MIN = 30


IMPRESSION_WORDS = ("impression", "покази", "показ", "перегляд", "reach", "охоплен")
ENGAGEMENT_WORDS = ("engagement", "взаємод", "залучен", "engagements")


def _metric_of_header(text):
    low = (text or "").strip().lower()
    if any(w in low for w in IMPRESSION_WORDS):
        return "reach"
    if any(w in low for w in ENGAGEMENT_WORDS):
        return "engagement"
    return None


NUMERIC_CELL = re.compile(r"^\d[\d\s, ]*$")


def _as_number(value):
    """The export writes every metric as text ('15921'), so a plain isinstance check
    finds nothing. Dates like '7/16/2026' must not pass — hence the digits-only rule."""
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return int(value)
    if isinstance(value, str) and NUMERIC_CELL.match(value.strip()):
        return int(re.sub(r"[\s, ]", "", value.strip()))
    return None


def from_metrics_xlsx(path):
    """AggregateAnalytics export (profile → Analytics → Export).

    Impressions are author-only data: no API and no scraper returns them, so this
    sheet is the single source for `reach`. The TOP POSTS sheet holds two top-50
    lists side by side — left keyed on engagements, right on impressions — so the
    same post appears in both blocks and is joined on the share id in its URL.

    Layout drifts between export vintages, so nothing here is hard-coded to a cell:
    every URL cell is paired with the first number to its right, and the metric name
    comes from the nearest header above that column pair.

    Each export counts impressions INSIDE its own date range only (a post from 14.08
    shows 11 232 in the August file and 165 in September's), so the window travels
    with the numbers: `lifetime_of` sums them over windows that do not overlap.

    Returns (window, {share_id: {"reach": int|None, "engagement": int|None}}),
    window = "YYYY-MM-DD_YYYY-MM-DD".
    """
    try:
        from openpyxl import load_workbook
    except ImportError:
        sys.exit("openpyxl is required for --metrics: pip3 install openpyxl")

    wb = load_workbook(path, data_only=True, read_only=True)
    sheets = [s for s in wb.sheetnames if "top" in s.lower() and "post" in s.lower()]
    metrics, matched = {}, []

    for name in sheets or wb.sheetnames:
        grid = [list(r) for r in wb[name].iter_rows(values_only=True)]
        # column -> metric, taken from any header cell naming a metric
        col_metric = {}
        for row in grid:
            for col, cell in enumerate(row):
                metric = _metric_of_header(cell if isinstance(cell, str) else "")
                if metric:
                    col_metric[col] = metric

        for row in grid:
            for col, cell in enumerate(row):
                if not isinstance(cell, str):
                    continue
                m = SHARE_ID.search(cell)
                if not m:
                    continue
                # first number to the right of the URL is that post's value
                for vcol in range(col + 1, len(row)):
                    value = _as_number(row[vcol])
                    if value is not None:
                        metric = col_metric.get(vcol) or col_metric.get(col)
                        if not metric:
                            break  # unlabelled column: better nothing than a guess
                        rec = metrics.setdefault(m.group(1), {})
                        rec[metric] = value
                        matched.append((name, metric))
                        break

    counts = {}
    for _, metric in matched:
        counts[metric] = counts.get(metric, 0) + 1
    window = window_of_export(wb, path)
    print(
        f"metrics: {len(metrics)} posts from {path.name} [{window}] "
        f"({counts or 'nothing matched'})"
    )
    if not metrics:
        print("  ! no post URLs paired with numbers — check the sheet layout:")
        print(f"    sheets: {wb.sheetnames}")
    return window, metrics


US_RANGE = re.compile(r"(\d{1,2})/(\d{1,2})/(\d{4})\s*-\s*(\d{1,2})/(\d{1,2})/(\d{4})")
ISO_RANGE = re.compile(r"(\d{4}-\d{2}-\d{2})_(\d{4}-\d{2}-\d{2})")


def window_of_export(wb, path):
    """Date range of an export: the DISCOVERY sheet states it ('7/1/2026 - 8/17/2026');
    the file name carries the same range and is the fallback."""
    for name in wb.sheetnames:
        if name.strip().upper() != "DISCOVERY":
            continue
        for row in wb[name].iter_rows(max_row=5, values_only=True):
            for cell in row:
                m = US_RANGE.search(str(cell or ""))
                if m:
                    m1, d1, y1, m2, d2, y2 = (int(x) for x in m.groups())
                    return f"{y1}-{m1:02d}-{d1:02d}_{y2}-{m2:02d}-{d2:02d}"
    m = ISO_RANGE.search(Path(path).name)
    if m:
        return f"{m.group(1)}_{m.group(2)}"
    sys.exit(f"{path}: no date range in DISCOVERY or the file name — cannot place its numbers")


def lifetime_of(windows, field):
    """Lifetime value of one metric = the largest sum over export windows that do not
    overlap (weighted interval scheduling).

    Disjoint windows simply add up. Overlapping ones (01.07–17.08 vs 01–31.08) cannot
    both count, so the better chain wins — a floor, never a double count. A post that
    missed a window's top-50 has no number there and adds nothing, also a floor.
    """
    spans = sorted(
        (tuple(key.split("_")), vals[field])
        for key, vals in (windows or {}).items()
        if vals.get(field) is not None
    )
    spans.sort(key=lambda span: span[0][1])
    if not spans:
        return None
    best = [0] * (len(spans) + 1)
    for k, ((start, _end), value) in enumerate(spans):
        prev = max((m + 1 for m in range(k) if spans[m][0][1] < start), default=0)
        best[k + 1] = max(best[k], best[prev] + value)
    return best[-1]


def take_windows(rec, observed):
    """Fold this run's per-window numbers into the record and recount its lifetime.

    Windows are kept on the record (`metrics_windows`), so a later run that passes
    only the newest export adds its tail instead of replacing the total with it."""
    windows = dict(rec.get("metrics_windows") or {})
    for window, vals in observed.items():
        windows[window] = {**windows.get(window, {}), **vals}
    rec["metrics_windows"] = dict(sorted(windows.items()))
    for field in ("reach", "engagement"):
        if not covers_launch(windows, rec.get("date"), field):
            # only later windows saw this post: what they hold is a tail without the
            # launch week, where most impressions land. A 37 for a post that did
            # thousands drags every median, so the lifetime stays unknown instead
            rec[field] = None
            continue
        value = lifetime_of(windows, field)
        if value is not None:
            rec[field] = value


def covers_launch(windows, published, field):
    """True when some export window holding this metric contains the publish date.
    One day of slack: the export stamps the date in GMT, so a late post lands a day
    later there than in the archive."""
    from datetime import date, timedelta

    try:
        day = date.fromisoformat(published)
    except (TypeError, ValueError):
        return True  # no date to judge by: trust the numbers
    for key, vals in windows.items():
        if vals.get(field) is None:
            continue
        start, end = (date.fromisoformat(x) for x in key.split("_"))
        if start <= day + timedelta(days=1) and day <= end:
            return True
    return False


def slugs_from_metrics_xlsx(path):
    """{share_id: {"url", "date", "slug"}} from the same TOP POSTS sheet.

    Needed because a Basic data export carries no Shares.csv: posts that exist only
    as local drafts have no url, so `apply_metrics` has no share id to join on. The
    post url embeds a slug of the opening line, which is enough to pair a draft with
    its analytics row.
    """
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True, read_only=True)
    out = {}
    for name in wb.sheetnames:
        if "top" not in name.lower() or "post" not in name.lower():
            continue
        for row in wb[name].iter_rows(values_only=True):
            for col, cell in enumerate(row):
                if not isinstance(cell, str) or "linkedin.com/posts/" not in cell:
                    continue
                m = SHARE_ID.search(cell)
                if not m:
                    continue
                slug = unquote(cell.rsplit("/posts/", 1)[-1].split("_", 1)[-1])
                slug = re.split(r"-(?:share|ugcPost)-", slug)[0]
                date = ""
                for vcol in range(col + 1, len(row)):
                    if isinstance(row[vcol], str) and "/" in row[vcol]:
                        mm, dd, yy = row[vcol].split("/")
                        date = f"{yy}-{int(mm):02d}-{int(dd):02d}"
                        break
                out[m.group(1)] = {"url": cell, "date": date, "slug": slug}
    return out


def _words(text):
    return [w for w in re.split(r"[^0-9a-zа-яіїєґ]+", (text or "").lower()) if w]


def apply_metrics_by_slug(records, metrics, slugs, used=()):
    """Pair url-less drafts with analytics rows on publish date + opening-line slug.

    LinkedIn stamps the publish date in GMT, so a late-evening post lands on the next
    day in the export — hence the +/-1 day window.
    """
    by_id = {}
    for rec in records.values():
        # a draft that kept its reach but lost its url (pre-fix rebuilds) still needs
        # pairing — skipping it is what seeded the phantom tail records
        if rec.get("url"):
            continue
        by_id[rec["id"]] = (rec, _words(rec.get("hook") or rec.get("text") or "")[:14])
    claimed = {m.group(1) for r in records.values() if (m := SHARE_ID.search(r.get("url") or ""))}

    hit = 0
    placed = set()
    for share_id, info in slugs.items():
        found = metrics.get(share_id)
        if not found or not info["date"] or share_id in used:
            continue
        head = _words(info["slug"])[:6]
        if len(head) < 3:
            continue
        best = None
        for rec, hook_words in by_id.values():
            if abs(_day_delta(rec.get("date"), info["date"])) > 1:
                continue
            overlap = len(set(head) & set(hook_words))
            if overlap >= 3 and (best is None or overlap > best[0]):
                best = (overlap, rec)
        if not best:
            continue
        rec = best[1]
        rec["url"] = info["url"]
        rec["open_url"] = open_url_of(info["url"])
        take_windows(rec, found)
        by_id.pop(rec["id"], None)
        hit += 1
        placed.add(share_id)

    # a post published straight from LinkedIn leaves no local draft: the export is
    # then the only record it has, so seed one from the slug rather than lose it
    added = 0
    for share_id, info in slugs.items():
        if share_id in placed or share_id in used or share_id in records or share_id in claimed:
            continue
        found = metrics.get(share_id)
        if not found or not info["date"]:
            continue
        hook = info["slug"].replace("-", " ").strip()
        records[share_id] = {
            "id": share_id,
            "date": info["date"],
            "hook": hook[:1].upper() + hook[1:],
            "url": info["url"],
            "open_url": open_url_of(info["url"]),
            "source": "metrics",
            "text": "",
        }
        take_windows(records[share_id], found)
        added += 1
    print(f"metrics matched by slug: {hit}; seeded from export: {added}")
    return records


def _day_delta(a, b):
    from datetime import date as _date

    try:
        pa = _date(*[int(x) for x in a.split("-")])
        pb = _date(*[int(x) for x in b.split("-")])
    except Exception:
        return 99
    return (pa - pb).days


def fold_seeded_into_drafts(records):
    """Heal records seeded from an export ("metrics") that are really a local draft's
    post: the draft lost its share id, so a later export's tail was seeded as a new
    post. Pair them on date +/-1 day and opening words, the same rule the slug pass
    uses, one-to-one by best overlap; the draft keeps its id and manual fields, the
    seeded record hands over its url and per-window numbers and is dropped."""
    seeded = [r for r in records.values() if r.get("source") == "metrics"]
    drafts = [r for r in records.values() if r.get("source") == "local-md" and not r.get("url")]
    pairs = []
    for s in seeded:
        head = _words(s.get("hook"))[:6]
        if len(head) < 3:
            continue
        for d in drafts:
            if abs(_day_delta(d.get("date"), s.get("date"))) > 1:
                continue
            overlap = len(set(head) & set(_words(d.get("hook") or d.get("text"))[:14]))
            if overlap >= 3:
                pairs.append((overlap, s["id"], d["id"]))
    taken_s, taken_d, folded = set(), set(), 0
    for overlap, sid, did in sorted(pairs, reverse=True):
        if sid in taken_s or did in taken_d:
            continue
        taken_s.add(sid)
        taken_d.add(did)
        s, d = records[sid], records[did]
        d["url"], d["open_url"] = s["url"], s.get("open_url") or open_url_of(s["url"])
        take_windows(d, s.get("metrics_windows") or {})
        for field in MANUAL_FIELDS:
            if s.get(field) and not d.get(field):
                d[field] = s[field]
        del records[sid]
        folded += 1
        print(f"  folded seeded {sid} ({s.get('date')}) into draft {did}")
    print(f"seeded records folded into their drafts: {folded}")
    return records


def apply_metrics(records, metrics):
    """Fill reach/engagement on records whose share id is in the analytics export.

    Returns (records, share ids that found a home) so the slug pass below does not
    hand the same numbers to a draft copy of a post that already got them.
    """
    hit = 0
    used = set()
    for rec in records.values():
        share_id = SHARE_ID.search(rec.get("url") or "")
        key = share_id.group(1) if share_id else rec["id"]
        found = metrics.get(key)
        if not found:
            continue
        hit += 1
        used.add(key)
        take_windows(rec, found)
    print(f"metrics applied to {hit} of {len(records)} posts")
    return records, used


STICKY_FIELDS = ("url", "open_url", "posted_at", "metrics_windows", "reach", "engagement")


def merge(existing, incoming):
    """Export text wins over local drafts; manual classification always survives."""
    for rec in incoming:
        old = existing.get(rec["id"])
        if old:
            for f in MANUAL_FIELDS:
                if old.get(f):
                    rec[f] = old[f]
            # a local draft re-read from disk has no url and no metrics; without this the
            # share id paired by slug last run is lost, the next export cannot join on it,
            # and its tail gets seeded as a second, phantom post
            for f in STICKY_FIELDS:
                if not rec.get(f) and old.get(f):
                    rec[f] = old[f]
            if old.get("source") == "linkedin-export" and rec["source"] == "local-md":
                rec["text"] = old["text"]
                rec["source"] = "linkedin-export"
                rec["url"] = old.get("url") or rec["url"]
        for f in MANUAL_FIELDS:
            rec.setdefault(f, "" if f != "topic_tags" else [])
        if not rec.get("open_url"):
            rec["open_url"] = (old or {}).get("open_url") or open_url_of(rec.get("url"))
        existing[rec["id"]] = rec
    return existing


def write_index(records):
    rows = sorted(records, key=lambda r: r.get("date") or "", reverse=True)
    lines = [
        "# ANGLES — індекс постів",
        "",
        f"Постів у архіві: **{len(rows)}**. Перебудова: `python3 <тека скіла 09-track>/scripts/build_archive.py --root .`.",
        "",
        "Читати ПЕРЕД кожним новим постом: шукати збіг теми + кута, щоб не писати дубль.",
        "",
        "| Дата | Хук | Кут | Рубрика | Теми | Охоплення |",
        "|---|---|---|---|---|---|",
    ]
    for r in rows:
        hook = (r.get("hook") or "").replace("|", "·")[:70]
        tags = " · ".join(r.get("topic_tags") or [])
        lines.append(
            f"| {r.get('date') or '?'} | {hook} | {r.get('angle') or ''} "
            f"| {r.get('rubric') or ''} | {tags} | {r.get('reach') or ''} |"
        )
    INDEX.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--shares", help="path to Shares.csv from the LinkedIn data export")
    ap.add_argument(
        "--metrics",
        type=Path,
        nargs="+",
        metavar="XLSX",
        help="AggregateAnalytics exports (profile → Analytics → Export). Pass several: "
        "each sheet only holds the top 50 posts of its own date range, so one file per "
        "quarter is what actually covers a year",
    )
    ap.add_argument("--media", help="path to Rich_Media.csv from the LinkedIn export")
    ap.add_argument(
        "--root",
        type=Path,
        default=Path.cwd(),
        help="author's working folder: drafts *.md are read from here, archive/ is written here "
        "(default: current directory)",
    )
    args = ap.parse_args()

    global ROOT, ARCHIVE, JSONL, INDEX
    ROOT = args.root.resolve()
    ARCHIVE = ROOT / "archive"
    JSONL = ARCHIVE / "posts.jsonl"
    INDEX = ARCHIVE / "ANGLES.md"

    ARCHIVE.mkdir(parents=True, exist_ok=True)
    records = load_existing()

    before = len(records)
    records = merge(records, from_local_md())
    if args.shares:
        records = merge(records, from_shares_csv(args.shares))
    # every run, not only with --shares: a draft written after the last export
    # otherwise sits beside its published post as a second record
    records = collapse_drafts(records)
    if args.media:
        records = apply_media(records, from_media_csv(args.media))
    records = apply_lead_magnets(records)
    if args.metrics:
        # {share_id: {window: {reach, engagement}}}; the same range exported twice
        # lands on one key, so a duplicate download never counts twice
        observed = {}
        for path in args.metrics:
            window, found_in = from_metrics_xlsx(path)
            for share_id, found in found_in.items():
                observed.setdefault(share_id, {})[window] = found
        records = fold_seeded_into_drafts(records)
        records, used = apply_metrics(records, observed)
        slugs = {}
        for path in args.metrics:
            slugs.update(slugs_from_metrics_xlsx(path))
        records = apply_metrics_by_slug(records, observed, slugs, used)

    ordered = sorted(records.values(), key=lambda r: r.get("date") or "", reverse=True)
    with open(JSONL, "w", encoding="utf-8") as fh:
        for rec in ordered:
            fh.write(json.dumps(rec, ensure_ascii=False) + "\n")
    write_index(ordered)

    unclassified = sum(1 for r in ordered if not r.get("angle"))
    print(f"posts: {len(ordered)} (+{len(ordered) - before} new)")
    print(f"unclassified angle: {unclassified}")
    print(f"-> {JSONL}")
    print(f"-> {INDEX}")


if __name__ == "__main__":
    main()
